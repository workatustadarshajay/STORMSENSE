"""Writes the four demo Excel workbooks for uploading: stores, products, daily sales and daily stock.

They come from the same generated sample data the planner uses, so the numbers are sample numbers, and the columns are
exactly the ones the upload expects. Sales and stock cover the last 28 days, which keeps each file well under the
5,000-row limit.

    python backend/scripts/make_demo_workbooks.py
"""

from __future__ import annotations

import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "databricks"))

from stormsense_core import reference, synth  # noqa: E402

END = date(2026, 10, 7)  # the same end date as the sample data
OUT = ROOT / "ingestion-client" / "public" / "demo"
DAYS = 28


def write(name: str, columns: list[str], rows: list[list]) -> None:
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = name
    ws.append(columns)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(14, len(col) + 4)
    book.save(OUT / f"{name}.xlsx")
    print(f"wrote {OUT / (name + '.xlsx')} ({len(rows)} rows)")


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    tables = synth.generate(END).tables
    since = END - timedelta(days=DAYS - 1)

    stores = reference.stores_df()
    write("stores", ["store_id", "name", "city", "region", "latitude", "longitude", "time_zone"],
          stores[["store_id", "name", "city", "region", "latitude", "longitude", "time_zone"]].values.tolist())

    products = reference.products_df()
    write("products", ["product_id", "name", "name_plural", "weather_driver", "unit_price", "pack_size"],
          products[["product_id", "name", "name_plural", "weather_driver", "unit_price", "pack_size"]].values.tolist())

    sales = tables["sales_history"]
    sales = sales[(sales["sale_date"] >= since) & (sales["sale_date"] <= END)]
    write("sales", ["store_id", "product_id", "sale_date", "units"],
          [[r.store_id, r.product_id, r.sale_date, int(r.units)] for r in sales.itertuples()])

    stock = tables["inventory_snapshot"]
    stock = stock[(stock["snapshot_date"] >= since) & (stock["snapshot_date"] <= END)]
    write("stock", ["store_id", "product_id", "snapshot_date", "on_hand", "in_transit"],
          [[r.store_id, r.product_id, r.snapshot_date, int(r.on_hand), int(r.in_transit or 0)] for r in stock.itertuples()])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
