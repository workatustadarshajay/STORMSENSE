"""Your own data, over HTTP: upload a CSV or send rows from another system. Off unless STORMSENSE_INGEST_ENABLED=1.

There is no sign-in on these routes (the hackathon copy is single-customer and local). Changes still need the
same-origin header, and they replace one feed at a time. Turn this off anywhere more than one person can reach.
"""

from __future__ import annotations

import base64
import csv
import io
import json
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, Request
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict, Field

from .auth import problem, same_origin
from .ingest import FEEDS, MAX_ROWS, IngestStore, auto_mapping, check, read_xlsx, template
from .service import Service


def _on(request: Request) -> None:
    if not request.app.state.settings.ingest_enabled:
        raise problem(404, "off", "Uploading your own data is turned off on this copy.")


router = APIRouter(prefix="/api/ingest", tags=["your data"], dependencies=[Depends(_on)])


class Refusal(BaseModel):
    line: int
    reason: str


class UploadResult(BaseModel):
    feed: str
    kept: int
    refused: int
    refusals: list[Refusal] = Field(description="The first 20 refused rows, with the line number and the reason")
    missing_columns: list[str]
    ignored_columns: list[str]
    message: str


class UploadBody(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    csv: str | None = Field(None, max_length=2_000_000, description="The whole CSV file as text, with the column names on the first line")
    xlsx_base64: str | None = Field(
        None, max_length=4_000_000, description="An Excel file (.xlsx), base64 encoded. Columns are matched by name"
    )
    mapping: dict[str, str] | None = Field(None, description="Your column name for each StormSense column, if the names differ")


class RowsBody(BaseModel):
    rows: list[dict[str, Any]] = Field(max_length=MAX_ROWS)
    mapping: dict[str, str] | None = None


class ColumnInfo(BaseModel):
    name: str
    kind: str
    required: bool
    help: str
    choices: list[str]


class FeedInfo(BaseModel):
    feed: str
    title: str
    what: str
    kept: int
    refused: int
    updated: str | None
    columns: list[ColumnInfo]
    template: str


def _feed(name: str):  # noqa: ANN202
    if name not in FEEDS:
        raise problem(404, "not_found", "There is no feed with that name.")
    return FEEDS[name]


def _result(feed, out, kept_now: int) -> UploadResult:  # noqa: ANN001
    if out.missing:
        message = "The file is missing columns: " + ", ".join(out.missing) + ". Nothing was kept."
    else:
        message = f"Kept {kept_now} rows. Refused {len(out.refused)}."
    return UploadResult(
        feed=feed.key,
        kept=kept_now,
        refused=len(out.refused),
        refusals=[Refusal(**r) for r in out.refused[:20]],
        missing_columns=out.missing,
        ignored_columns=out.ignored,
        message=message,
    )


def _store(request: Request):  # noqa: ANN202
    return request.app.state.ingest


@router.get("/feeds", response_model=list[FeedInfo])
def feeds(request: Request) -> list[FeedInfo]:
    store = _store(request)
    status = {s["feed"]: s for s in store.status()}
    return [
        FeedInfo(
            feed=f.key,
            title=f.title,
            what=f.what,
            kept=status[f.key]["rows"],
            refused=status[f.key]["refused"],
            updated=status[f.key]["updated"],
            columns=[ColumnInfo(name=c.name, kind=c.kind, required=c.required, help=c.help, choices=list(c.choices)) for c in f.columns],
            template=template(f),
        )
        for f in FEEDS.values()
    ]


@router.post("/feeds/{feed}/upload", response_model=UploadResult, dependencies=[Depends(same_origin)])
def upload(feed: str, body: UploadBody, request: Request) -> UploadResult:
    f = _feed(feed)
    if (body.csv is None) == (body.xlsx_base64 is None):
        raise problem(422, "one_file", "Send one file: a CSV, or an Excel workbook.")
    if body.xlsx_base64 is not None:
        rows = _excel_rows(f.key, body.xlsx_base64)
    else:
        rows = list(csv.DictReader(io.StringIO(body.csv)))
    if len(rows) > MAX_ROWS:
        raise problem(422, "too_many", f"A file can have at most {MAX_ROWS} rows. This one has {len(rows)}.")
    return _apply(f, rows, body.mapping, request)


@router.post("/feeds/{feed}/rows", response_model=UploadResult, dependencies=[Depends(same_origin)])
def rows_in(feed: str, body: RowsBody, request: Request) -> UploadResult:
    """Send rows from another system as JSON. Same checks as a file upload."""
    f = _feed(feed)
    rows = [{k: "" if v is None else str(v) for k, v in r.items()} for r in body.rows]
    return _apply(f, rows, body.mapping, request)


def _excel_rows(feed: str, encoded: str) -> list[dict[str, str]]:
    import base64
    import binascii
    import zipfile


    try:
        return read_xlsx(base64.b64decode(encoded, validate=True), feed)
    except (binascii.Error, zipfile.BadZipFile, KeyError, ValueError):
        raise problem(422, "not_excel", "That Excel file could not be read. Save it as .xlsx and try again.") from None


def _log(store: IngestStore, kind: str, **detail) -> None:  # noqa: ANN003
    from .events import record

    record(store.folder, kind, "upload app", **detail)


def _apply(f, rows: list[dict[str, str]], mapping: dict[str, str] | None, request: Request) -> UploadResult:  # noqa: ANN001
    store = _store(request)
    try:
        out = check(f, rows, mapping, store.known())
    except ValueError as e:  # personal or card columns
        raise problem(422, "not_accepted", str(e)) from None
    if not out.missing:
        store.save(f.key, out, {**auto_mapping(f, list(rows[0].keys()) if rows else []), **(mapping or {})})
        _log(store, "upload", feed=f.title, kept=len(out.accepted), refused=len(out.refused))
    return _result(f, out, len(out.accepted) if not out.missing else 0)


class AnalysisItem(BaseModel):
    store: str
    product: str
    on_hand: int
    sold_per_day: float
    days_of_cover: float | None
    status: str = Field(description="Short (under 7 days), Watch (under 21 days), Plenty, or No sales")
    stock_value_usd: float
    sales_value_usd: float


class AnalysisStore(BaseModel):
    store: str
    short_items: int
    stock_value_usd: float


class Analysis(BaseModel):
    as_of: str
    window_days: int
    stores: int
    products: int
    pairs: int
    short: int
    watch: int
    sold_units: int
    sales_value_usd: float
    stock_value_usd: float
    by_store: list[AnalysisStore]
    items: list[AnalysisItem] = Field(description="The short and watch items, most urgent first")


@router.get("/analysis", response_model=Analysis)
def analysis(request: Request) -> Analysis:
    """A plain analysis of your uploaded data. Needs the stores, products, sales and stock files."""
    from .ingest_analysis import analyse

    result = analyse(_store(request))
    if result is None:
        raise problem(404, "not_ready", "Upload your stores, products, daily sales and daily stock files first.")
    return Analysis(**result)


@router.get("/feeds/{feed}/preview")
def preview(feed: str, request: Request, limit: int = 20) -> list[dict[str, Any]]:
    _feed(feed)
    saved = _store(request).read(feed)
    return saved["rows"][: max(1, min(limit, 100))] if saved else []


@router.delete("/feeds/{feed}", dependencies=[Depends(same_origin)])
def clear(feed: str, request: Request) -> dict[str, str]:
    _feed(feed)
    _store(request).clear(feed)
    return {"message": "Cleared. The planning screens will use the sample data for this file again."}


# ---- checks, plan, economics, demo, templates ------------------------------------------------

REPO = Path(__file__).resolve().parents[2]
DEMO_DIR = REPO / "ingestion-client" / "public" / "demo"


class ChecksResult(BaseModel):
    sentences: list[str]


class PlanStatus(BaseModel):
    ready: bool
    as_of: str | None = None
    summary: dict[str, Any] | None = None
    net_benefit: dict[str, float] | None = None


class EconomicsBody(BaseModel):
    truck_cost_per_mile: float = Field(ge=0, le=20, description="What one truck trip costs per mile, in dollars")
    margin_pct: float = Field(ge=0, le=100, description="Share of sales that is profit, in percent")


def _plan_status(store: IngestStore) -> PlanStatus:
    from .ingest_plan import PLAN_FILE, economics, net_benefit

    path = store.folder / PLAN_FILE
    if not path.exists():
        return PlanStatus(ready=False)
    plan = json.loads(path.read_text())
    return PlanStatus(ready=True, as_of=plan["as_of"], summary=plan["summary"], net_benefit=net_benefit(plan["summary"], economics(store)))


@router.get("/checks", response_model=ChecksResult)
def upload_checks(request: Request) -> ChecksResult:
    from .ingest_analysis import checks

    return ChecksResult(sentences=checks(_store(request)))


class DropResult(BaseModel):
    folder: str
    loaded: list[str]


@router.post("/drop/scan", response_model=DropResult, dependencies=[Depends(same_origin)])
def scan_drop_now(request: Request) -> DropResult:
    """Loads any files waiting in the drop folder now, instead of waiting for the next check."""
    from .ingest import DROP_DIR, scan_drop

    store = _store(request)
    loaded = scan_drop(store)
    if loaded:
        _log(store, "drop_loaded", feeds=loaded)
    return DropResult(folder=str(store.folder / DROP_DIR), loaded=loaded)


@router.get("/charts")
def upload_charts_route(request: Request) -> dict[str, Any]:
    """Chart data about your uploads: sales and stock by day, cover by store, status counts, and what was loaded."""
    from .ingest_charts import upload_charts

    return upload_charts(_store(request))


@router.get("/plan/status", response_model=PlanStatus)
def plan_status(request: Request) -> PlanStatus:
    return _plan_status(_store(request))


@router.post("/plan", response_model=PlanStatus, dependencies=[Depends(same_origin)])
def build_plan(request: Request) -> PlanStatus:
    """Builds the plan from your uploads. The planner's screens can then use them as "Your uploads"."""
    from .ingest_plan import PlanNotReady, build, save
    from .sources.mock import MockSource

    store = _store(request)
    try:
        plan = build(store)
    except PlanNotReady as e:
        raise problem(409, "not_ready", str(e)) from None
    path = save(store, plan)
    request.app.state.services["upload"] = Service(MockSource(path), request.app.state.settings)
    _log(store, "plan_built", transfers=plan["summary"]["transfers"], short=plan["summary"]["short"], as_of=plan["as_of"])
    return _plan_status(store)


@router.get("/economics", response_model=EconomicsBody)
def get_economics(request: Request) -> EconomicsBody:
    from .ingest_plan import economics

    return EconomicsBody(**economics(_store(request)))


@router.put("/economics", response_model=EconomicsBody, dependencies=[Depends(same_origin)])
def set_economics(body: EconomicsBody, request: Request) -> EconomicsBody:
    """Truck cost per mile and product margin, used for the estimated profit of the plan."""
    store = _store(request)
    (store.folder / "economics.json").write_text(json.dumps(body.model_dump()))
    _log(store, "economics_saved", truck_cost_per_mile=body.truck_cost_per_mile, margin_pct=body.margin_pct)
    return body


@router.post("/demo/load", dependencies=[Depends(same_origin)])
def load_demo(request: Request) -> dict[str, Any]:
    """Loads the four sample workbooks in order, then builds the plan. One click for a demo."""
    from .ingest_plan import PlanNotReady, build, save
    from .sources.mock import MockSource

    store = _store(request)
    loaded = {}
    for feed in ("stores", "products", "sales", "stock"):
        path = DEMO_DIR / f"{feed}.xlsx"
        if not path.exists():
            raise problem(404, "no_demo", "The sample workbooks are missing from this copy.")
        out = check(FEEDS[feed], _excel_rows(feed, base64.b64encode(path.read_bytes()).decode()), None, store.known())
        if out.missing:
            raise problem(422, "not_accepted", f"The sample {feed} workbook is missing columns: {', '.join(out.missing)}.")
        store.save(feed, out, auto_mapping(FEEDS[feed], [c.name for c in FEEDS[feed].columns]))
        loaded[feed] = len(out.accepted)
    try:
        plan = build(store)
    except PlanNotReady as e:  # cannot happen with the sample files, but keep the message plain
        raise problem(409, "not_ready", str(e)) from None
    request.app.state.services["upload"] = Service(MockSource(save(store, plan)), request.app.state.settings)
    _log(store, "demo_loaded", feeds=list(loaded))
    _log(store, "plan_built", transfers=plan["summary"]["transfers"], short=plan["summary"]["short"], as_of=plan["as_of"])
    return {"loaded": loaded, "checks": checks_for(store), "plan": _plan_status(store).model_dump()}


def checks_for(store: IngestStore) -> list[str]:
    from .ingest_analysis import checks

    return checks(store)


@router.get("/feeds/{feed}/template.xlsx")
def template_xlsx(feed: str, request: Request) -> Response:
    """An Excel template. Store and product codes are dropdowns, filled from the stores and products you uploaded."""
    f = _feed(feed)
    store = _store(request)
    body = template_workbook(f, _codes(store, "stores", "store_id"), _codes(store, "products", "product_id"))
    return Response(
        body,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{feed}-template.xlsx"'},
    )


def _codes(store: IngestStore, feed: str, field_name: str) -> list[str]:
    saved = store.read(feed)
    return [str(r[field_name]) for r in saved["rows"]] if saved else []


def template_workbook(feed, store_codes: list[str], product_codes: list[str]) -> bytes:  # noqa: ANN001
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    book = openpyxl.Workbook()
    ws = book.active
    ws.title = feed.key
    ws.append([c.name for c in feed.columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    lists = book.create_sheet("lists")
    lists["A1"], lists["B1"] = "store_id", "product_id"
    for i, code in enumerate(store_codes, start=2):
        lists.cell(i, 1, code)
    for i, code in enumerate(product_codes, start=2):
        lists.cell(i, 2, code)
    lists.sheet_state = "hidden"
    names = [c.name for c in feed.columns]
    for name, codes, col_letter in (("store_id", store_codes, "A"), ("product_id", product_codes, "B")):
        if name not in names or not codes:
            continue
        letter = get_column_letter(names.index(name) + 1)
        dv = DataValidation(
            type="list",
            formula1=f"=lists!${col_letter}$2:${col_letter}${len(codes) + 1}",
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Unknown code",
            error="Pick a code from the list. Upload the stores or products file first, then download this template again.",
        )
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{MAX_ROWS + 1}")
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()
