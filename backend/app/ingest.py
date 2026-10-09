"""Your own data: the four files the planning app reads, checked and stored locally.

Feeds: stores, products, sales (daily units sold) and stock (daily units on hand). Each upload replaces that feed.
Every row is checked before it is kept; refused rows come back with the line number and the reason.
A column whose name looks like personal or card data is refused outright.
"""

from __future__ import annotations

import csv
import io
import json
import re
import threading
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

MAX_ROWS = 5000
SENSITIVE_TOKENS = {"ssn", "aadhaar", "aadhar", "passport", "card", "credit", "patient", "dob", "birth", "license", "licence", "pan"}


@dataclass(frozen=True)
class Column:
    name: str
    kind: str  # text, number, whole, date
    required: bool = True
    minimum: float | None = None
    maximum: float | None = None
    choices: tuple[str, ...] = ()
    help: str = ""


@dataclass(frozen=True)
class Feed:
    key: str
    title: str
    what: str
    columns: tuple[Column, ...]
    unique: tuple[str, ...]


FEEDS: dict[str, Feed] = {
    "stores": Feed("stores", "Stores", "One row per store.", (
        Column("store_id", "text", help="Your store code, for example S01"),
        Column("name", "text"), Column("city", "text"), Column("region", "text", help="For example Florida"),
        Column("latitude", "number", minimum=-90, maximum=90), Column("longitude", "number", minimum=-180, maximum=180),
        Column("time_zone", "text", required=False, help="For example America/New_York"),
    ), unique=("store_id",)),
    "products": Feed("products", "Products", "One row per product.", (
        Column("product_id", "text", help="Your product code, for example P01"),
        Column("name", "text", help="Singular name, for example 1000W generator"),
        Column("name_plural", "text", help="Plural name, for example 1000W generators"),
        Column("weather_driver", "text", choices=("wind", "rain", "heat"), help="What weather moves its sales most"),
        Column("unit_price", "number", minimum=0.01),
        Column("pack_size", "whole", minimum=1, help="Units in one pack; moves are made in whole packs"),
    ), unique=("product_id",)),
    "sales": Feed("sales", "Daily sales", "Units sold per store, product and day.", (
        Column("store_id", "text"), Column("product_id", "text"),
        Column("sale_date", "date", help="For example 2026-10-01"),
        Column("units", "whole", minimum=0),
    ), unique=("store_id", "product_id", "sale_date")),
    "stock": Feed("stock", "Daily stock", "Units on hand per store, product and day.", (
        Column("store_id", "text"), Column("product_id", "text"),
        Column("snapshot_date", "date", help="The day the count was taken"),
        Column("on_hand", "whole", minimum=0),
        Column("in_transit", "whole", required=False, minimum=0, help="Units on the way; leave blank for zero"),
    ), unique=("store_id", "product_id", "snapshot_date")),
}
# Rows that point at a store or product must match a row in those feeds.
NEEDS = {"sales": ("stores", "products"), "stock": ("stores", "products")}


def _tokens(name: str) -> set[str]:
    return set(re.split(r"[^a-z]+", name.lower())) - {""}


def sensitive_columns(names: list[str]) -> list[str]:
    return [n for n in names if _tokens(n) & SENSITIVE_TOKENS]


def template(feed: Feed) -> str:
    buf = io.StringIO()
    csv.writer(buf).writerow([c.name for c in feed.columns])
    return buf.getvalue()


def _norm(s: str) -> str:
    """Column names match ignoring case, spaces and underscores, so "Store ID" finds store_id."""
    return re.sub(r"[\s_]+", " ", s.strip().lower())


def auto_mapping(feed: Feed, their_columns: list[str]) -> dict[str, str]:
    """Pair each StormSense column with a column in the file, matching names loosely."""
    lookup = {_norm(c): c for c in their_columns}
    return {c.name: lookup[_norm(c.name)] for c in feed.columns if _norm(c.name) in lookup}


@dataclass
class Outcome:
    accepted: list[dict[str, Any]] = field(default_factory=list)
    refused: list[dict[str, Any]] = field(default_factory=list)  # {line, reason}
    missing: list[str] = field(default_factory=list)
    ignored: list[str] = field(default_factory=list)


def _parse(col: Column, raw: Any) -> Any:
    text = "" if raw is None else str(raw).strip()
    if text == "":
        if col.required:
            raise ValueError(f"{col.name} is empty")
        return 0 if col.kind == "whole" else None
    if col.kind == "text":
        if col.choices and text.lower() not in col.choices:
            raise ValueError(f"{col.name} must be one of {', '.join(col.choices)}")
        return text.lower() if col.choices else text
    if col.kind == "date":
        try:
            return datetime.strptime(text[:10], "%Y-%m-%d").date().isoformat()
        except ValueError:
            raise ValueError(f"{col.name} must be a date like 2026-10-01") from None
    try:
        value = float(text.replace(",", ""))
    except ValueError:
        raise ValueError(f"{col.name} must be a number") from None
    if col.kind == "whole":
        if value != int(value):
            raise ValueError(f"{col.name} must be a whole number")
        value = int(value)
    if col.minimum is not None and value < col.minimum:
        raise ValueError(f"{col.name} must be at least {col.minimum:g}")
    if col.maximum is not None and value > col.maximum:
        raise ValueError(f"{col.name} must be at most {col.maximum:g}")
    return value


def check(feed: Feed, rows: list[dict[str, Any]], mapping: dict[str, str] | None,
          known: dict[str, set[str]] | None = None) -> Outcome:
    """Checks one upload. `rows` are dicts keyed by the file's own column names. `known` holds the codes already loaded."""
    out = Outcome()
    their = list(rows[0].keys()) if rows else []
    bad = sensitive_columns(their)
    if bad:
        raise ValueError("These columns look like personal or card data and are not accepted: " + ", ".join(bad))
    mapping = {**auto_mapping(feed, their), **(mapping or {})}
    out.missing = [c.name for c in feed.columns if c.required and c.name not in mapping]
    out.ignored = [c for c in their if c not in mapping.values()]
    if out.missing:
        return out
    seen: set[tuple] = set()
    for line, row in enumerate(rows, start=2):  # line 1 is the header
        try:
            clean = {c.name: _parse(c, row.get(mapping[c.name]) if c.name in mapping else None) for c in feed.columns}
            for ref in NEEDS.get(feed.key, ()):
                code_col = "store_id" if ref == "stores" else "product_id"
                if known is None or ref not in known:
                    raise ValueError(f"Upload the {ref} file first")
                if clean[code_col] not in known[ref]:
                    raise ValueError(f"{code_col} {clean[code_col]} is not in your {ref} file")
            key = tuple(clean[k] for k in feed.unique)
            if key in seen:
                raise ValueError("this row repeats an earlier one for the same key")
            seen.add(key)
            out.accepted.append(clean)
        except ValueError as e:
            out.refused.append({"line": line, "reason": str(e)})
    return out


class IngestStore:
    """Keeps each feed as a JSON file under `folder`. One lock, so two uploads cannot interleave."""

    def __init__(self, folder: Path) -> None:
        self.folder = folder
        self.folder.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()

    def _path(self, feed: str) -> Path:
        return self.folder / f"{feed}.json"

    def read(self, feed: str) -> dict[str, Any] | None:
        p = self._path(feed)
        return json.loads(p.read_text()) if p.exists() else None

    def codes(self, feed: str, field_name: str) -> set[str]:
        saved = self.read(feed)
        return {r[field_name] for r in saved["rows"]} if saved else set()

    def known(self) -> dict[str, set[str]]:
        """Codes already loaded. A feed that was never uploaded is left out, so the check can say to upload it first."""
        found = {}
        for feed, field_name in (("stores", "store_id"), ("products", "product_id")):
            if self.read(feed) is not None:
                found[feed] = self.codes(feed, field_name)
        return found

    def save(self, feed: str, outcome: Outcome, mapping: dict[str, str], replace: bool = True) -> None:
        with self._lock:
            self._path(feed).write_text(json.dumps({
                "rows": outcome.accepted, "mapping": mapping, "refused": len(outcome.refused),
                "updated": datetime.now().isoformat(timespec="seconds"),
            }, default=str))

    def clear(self, feed: str) -> None:
        with self._lock:
            self._path(feed).unlink(missing_ok=True)

    def status(self) -> list[dict[str, Any]]:
        out = []
        for f in FEEDS.values():
            saved = self.read(f.key)
            out.append({"feed": f.key, "title": f.title, "rows": len(saved["rows"]) if saved else 0,
                        "refused": saved["refused"] if saved else 0, "updated": saved["updated"] if saved else None})
        return out


def read_xlsx(data: bytes, feed_key: str) -> list[dict[str, str]]:
    """Rows of an Excel workbook, as text keyed by the header row. Uses the sheet named after the feed, else the first sheet."""
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = book[feed_key] if feed_key in book.sheetnames else book.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return []
        names = [str(h).strip() if h is not None else "" for h in header]
        out = []
        for values in rows:
            if all(v is None or str(v).strip() == "" for v in values):
                continue  # blank lines in a spreadsheet are not rows
            row = {}
            for name, v in zip(names, values):
                if not name:
                    continue
                if isinstance(v, datetime):
                    row[name] = v.date().isoformat()
                elif isinstance(v, date):
                    row[name] = v.isoformat()
                else:
                    row[name] = "" if v is None else str(v)
            out.append(row)
        return out
    finally:
        book.close()
