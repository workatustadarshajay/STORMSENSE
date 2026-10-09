#!/usr/bin/env python3
"""Second opinion on the deck: open it with python-pptx (an independent OOXML
implementation) and re-derive the facts the builder claims.

Run:  PYTHONPATH=/tmp/pptxlib python3 deck/verify_with_pptx.py
"""

import os
import sys

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:                                         # pragma: no cover
    print("python-pptx is not installed - cross-check skipped "
          "(pip install python-pptx)")
    raise SystemExit(0)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DECK = os.path.join(ROOT, "StormSense-Databricks-Hackathon.pptx")

fails = []
prs = Presentation(DECK)

if (prs.slide_width, prs.slide_height) != (12192000, 6858000):
    fails.append(f"slide size {prs.slide_width}x{prs.slide_height}")
if len(prs.slides) != 22:
    fails.append(f"{len(prs.slides)} slides, expected 22")

titles, notes, media, tables = [], 0, 0, 0
for i, slide in enumerate(prs.slides, 1):
    t = ""
    try:
        if slide.shapes.title is not None:
            t = slide.shapes.title.text_frame.text.strip()
    except Exception:                                       # noqa: BLE001
        t = ""
    titles.append(t)
    for shape in slide.shapes:
        if shape.shape_type == 16 or shape.has_table if False else False:
            pass
        if getattr(shape, "has_table", False) and shape.has_table:
            tables += 1
            rows = [[c.text.strip() for c in r.cells] for r in shape.table.rows]
            if len(rows) != 7 or len(rows[0]) != 3:
                fails.append(f"slide {i}: table shape {len(rows)}x{len(rows[0])}")
        if shape.shape_type == 16:                          # MEDIA
            media += 1
        left, top = Emu(shape.left), Emu(shape.top)
        right = shape.left + shape.width
        bottom = shape.top + shape.height
        if shape.left < 0 or shape.top < 0 or right > 12192000 or bottom > 6858000:
            fails.append(f"slide {i}: {shape.name} outside the slide")
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
        notes += 1

print(f"slides           : {len(prs.slides)}")
print(f"slide size       : {prs.slide_width} x {prs.slide_height} EMU")
print(f"speaker notes on : {notes} slides (slide 3 keeps the template's empty one)")
print(f"tables           : {tables}")
print(f"media shapes     : {media}")
print()
for i, t in enumerate(titles, 1):
    print(f"{i:>2}. {t}")

if media != 1:
    fails.append(f"{media} media shapes, expected 1")
# 12 written slides carry notes; the template's own slide 3 has an empty notes
# part, so 12 is the number of slides with real speaker notes.
if notes != 12:
    fails.append(f"notes on {notes} slides, expected 12")
if tables != 1:
    fails.append(f"{tables} tables, expected 1")

print()
if fails:
    for f in fails:
        print("FAIL", f)
    sys.exit(1)
print("python-pptx cross-check: all good")
